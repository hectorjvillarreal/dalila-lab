"""Genera los dos .docx de la sección 10 con el formato del documento de Juan Pablo (lo leído en su XML):

- `N13_seccion10_esquemas_financiamiento_v1.docx`: la sección sola, sobre la plantilla de su documento (estilos, márgenes,
  pie de página), sin su texto ni sus comentarios;
- `Nucleo13_Informe_Resultados_Oct2026 atendido JP_con_seccion10_v1.docx`: copia de su documento con la sección al final,
  después de «9. Anexos» (opción b), en página nueva.

Formato copiado del documento: títulos con los estilos «Heading 1» y «Heading 2»; párrafos justificados con 6 pt antes y
después e interlineado 1.15; figura centrada con la leyenda debajo (negrita, 9 pt) y la fuente al pie (cursiva gris, 8 pt);
título de tabla arriba (negrita, 9 pt); tablas a todo el ancho de texto, encabezado con fondo D9E2F3, bordes BFBFBF y
texto de 8.5 pt. El documento no usa listas numeradas: las viñetas se escriben como párrafos con sangría francesa.
Las marcas [REVISAR…] se resaltan en amarillo. El original no se modifica; el script verifica su hash.

Desde la v1.1: la versión se pasa como argumento (`python docx_seccion10.py v1.1`, por omisión la v1.1); figuras al 85 % del
ancho de texto (5.4 in); filas de tabla que no se parten entre páginas y encabezado repetido; notas al pie del .md
(`[^n]` y `[^n]: texto`) como notas de Word con los estilos «FootnoteText» y «FootnoteReference» de JP."""

from __future__ import annotations

import hashlib
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

RAIZ = Path(__file__).resolve().parents[1]
VERSION = sys.argv[1] if len(sys.argv) > 1 else "v1.1"
ORIGINAL = RAIZ.parent / "Nucleo13_Informe_Resultados_Oct2026 atendido JP.docx"
FUENTE_MD = RAIZ / f"N13_seccion10_esquemas_financiamiento_{VERSION}.md"
SALIDA_SOLA = RAIZ / f"N13_seccion10_esquemas_financiamiento_{VERSION}.docx"
SALIDA_COPIA = RAIZ / f"Nucleo13_Informe_Resultados_Oct2026 atendido JP_con_seccion10_{VERSION}.docx"
ANCHO_FIG = 5.4 if VERSION != "v1" else 6.0     # pulgadas
ANCHO_TEXTO = 9180          # twips: carta (12240) menos márgenes 1800 y 1260
GRIS = RGBColor(0x59, 0x59, 0x59)
MARGEN_V = 100 if VERSION == "v1" else 50     # margen superior e inferior de celda (twips); JP usa 100, la v1.1 compacta las tablas


# ---------------------------------------------------------------- lectura del .md

NOTAS: dict[str, str] = {}      # texto de cada nota al pie del .md ([^n]: …)


def bloques(md: str) -> list[tuple[str, object]]:
    """Divide el .md en bloques: h1, h2, p, li, img, fig (leyenda), tab (título), fuente, tabla. Guarda las notas al pie."""
    NOTAS.clear()
    for m in re.finditer(r"^\[\^(\d+)\]: (.+)$", md, flags=re.M):
        NOTAS[m.group(1)] = m.group(2).strip()
    md = re.sub(r"^\[\^\d+\]: .+$", "", md, flags=re.M)
    out, lineas, i = [], md.splitlines(), 0
    while i < len(lineas):
        ln = lineas[i].rstrip()
        if not ln:
            i += 1
            continue
        if ln.startswith("# "):
            out.append(("h1", ln[2:]))
        elif ln.startswith("## "):
            out.append(("h2", ln[3:]))
        elif ln.startswith("!["):
            out.append(("img", re.search(r"\(([^)]+)\)", ln).group(1)))
        elif ln.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                if not re.fullmatch(r"\|[\s\-:|]+\|", lineas[i].strip()):
                    filas.append([c.strip() for c in lineas[i].strip().strip("|").split("|")])
                i += 1
            out.append(("tabla", filas))
            continue
        elif ln.startswith("- "):
            out.append(("li", ln[2:]))
        elif re.fullmatch(r"\*\*Figura 10\.\d+\..*\*\*", ln):
            out.append(("fig", ln[2:-2]))
        elif re.fullmatch(r"\*\*Tabla 10\.\d+\..*\*\*", ln):
            out.append(("tab", ln[2:-2]))
        elif re.fullmatch(r"\*Fuente:.*\*", ln):
            out.append(("fuente", ln[1:-1]))
        else:
            out.append(("p", ln))
        i += 1
    return out


def corridas(texto: str) -> list[tuple[str, bool, bool]]:
    """Separa **negrita**, *cursiva* y texto normal."""
    partes = []
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*|([^*]+)", texto):
        if m.group(1) is not None:
            partes.append((m.group(1), True, False))
        elif m.group(2) is not None:
            partes.append((m.group(2), False, True))
        else:
            partes.append((m.group(3), False, False))
    return partes


# ---------------------------------------------------------------- formato

def _espaciado(p, antes, despues, jc=None):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Twips(antes), Twips(despues)
    pf.line_spacing = 1.15
    if jc:
        p.alignment = jc


class _Notas:
    """Notas al pie de Word sobre la parte footnotes.xml del documento (python-docx no la modela)."""
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    def __init__(self, doc, limpiar: bool):
        self.parte = next(r.target_part for r in doc.part.rels.values() if r.reltype.endswith("/footnotes"))
        xml = self.parte.blob.decode("utf-8")
        if limpiar:     # plantilla sin el texto de JP: quitar sus notas (quedan los separadores -1 y 0)
            xml = re.sub(r'<w:footnote w:id="[1-9]\d*">.*?</w:footnote>', "", xml, flags=re.S)
        self.xml = xml
        ids = [int(x) for x in re.findall(r'<w:footnote\b[^>]*\bw:id="(-?\d+)"', xml)]
        self.sig = max(ids + [0]) + 1

    def agregar(self, texto: str) -> int:
        n = self.sig
        self.sig += 1
        t = texto.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        self.xml = self.xml.replace("</w:footnotes>",
            f'<w:footnote w:id="{n}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
            f'<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteRef/></w:r>'
            f'<w:r><w:t xml:space="preserve"> {t}</w:t></w:r></w:p></w:footnote></w:footnotes>')
        return n

    def guardar(self):
        self.parte._blob = self.xml.encode("utf-8")


NOTA_DOC: list = [None]


def _llamada(p, n_md: str):
    n = NOTA_DOC[0].agregar(NOTAS[n_md])
    r = p.add_run()
    rPr = r._r.get_or_add_rPr()
    st = OxmlElement("w:rStyle")
    st.set(qn("w:val"), "FootnoteReference")
    rPr.append(st)
    fr = OxmlElement("w:footnoteReference")
    fr.set(qn("w:id"), str(n))
    r._r.append(fr)


def _texto(p, texto, tam=None, negrita=False, cursiva=False, color=None):
    for trozo in re.split(r"(\[REVISAR[^\]]*\]|\[\^\d+\])", texto):
        if not trozo:
            continue
        if re.fullmatch(r"\[\^\d+\]", trozo):
            _llamada(p, trozo[2:-1])
            continue
        for t, b, it in corridas(trozo):
            r = p.add_run(t)
            r.bold = negrita or b or None
            r.italic = cursiva or it or None
            if tam:
                r.font.size = Pt(tam)
            if color:
                r.font.color.rgb = color
            if trozo.startswith("[REVISAR"):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _celda(celda, texto, encabezado, ancho):
    tcPr = celda._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(ancho))
    tcW.set(qn("w:type"), "dxa")
    bordes = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single"), b.set(qn("w:sz"), "6"), b.set(qn("w:space"), "0"), b.set(qn("w:color"), "BFBFBF")
        bordes.append(b)
    tcPr.append(bordes)
    if encabezado:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), "D9E2F3")
        tcPr.append(shd)
    mar = OxmlElement("w:tcMar")
    for lado, v in (("top", MARGEN_V), ("left", 80), ("bottom", MARGEN_V), ("right", 80)):
        m = OxmlElement(f"w:{lado}")
        m.set(qn("w:w"), str(v)), m.set(qn("w:type"), "dxa")
        mar.append(m)
    tcPr.append(mar)
    p = celda.paragraphs[0]
    _espaciado(p, 20, 20)
    p.alignment = None
    _texto(p, texto, tam=8.5, negrita=encabezado)


def _anchos_min_altura(filas):
    """Estima los renglones de cada fila (8.5 pt ≈ 86 twips por carácter, menos 160 de márgenes) y elige los anchos que
    minimizan la suma de renglones de la tabla."""
    import itertools
    import math
    n = len(filas[0])
    pasos = [x / 40 for x in range(4, 41)]
    mejor, best = None, None
    for comb in itertools.product(pasos, repeat=n - 1):
        ult = 1 - sum(comb)
        if ult < 0.10 - 1e-9:
            continue
        fr = list(comb) + [ult]
        cpl = [max(1, (ANCHO_TEXTO * f - 160) / 86) for f in fr]
        alto = sum(max(math.ceil(len(c) / k) for c, k in zip(fila, cpl)) for fila in filas)
        if best is None or alto < best:
            mejor, best = fr, alto
    anchos = [round(ANCHO_TEXTO * f) for f in mejor]
    anchos[-1] += ANCHO_TEXTO - sum(anchos)
    return anchos


def _anchos(filas):
    """Ancho de columna proporcional al texto más largo, con mínimo, sumando el ancho de texto."""
    n = len(filas[0])
    if VERSION == "v1":
        largo = [max(min(len(f[j]), 60) for f in filas) for j in range(n)]
        base = [max(l, 10) for l in largo]
    else:   # v1.1: la combinación de anchos (pasos de 2.5 %, mínimo 10 %) que da la tabla más baja
        return _anchos_min_altura(filas)
    anchos = [round(ANCHO_TEXTO * b / sum(base)) for b in base]
    anchos[-1] += ANCHO_TEXTO - sum(anchos)
    return anchos


def _tabla(doc, filas):
    anchos = _anchos(filas)
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    tblPr = t._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        e = tblPr.find(qn(tag))
        if e is not None:
            tblPr.remove(e)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(ANCHO_TEXTO)), tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), anchos):
        gc.set(qn("w:w"), str(w))
    for i, fila in enumerate(filas):
        tr = t.rows[i]._tr
        trPr = tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))      # la fila no se parte entre páginas
        if i == 0:   # repetir el encabezado si la tabla cruza de página
            trPr.append(OxmlElement("w:tblHeader"))
        for j, texto in enumerate(fila):
            _celda(t.cell(i, j), texto, i == 0, anchos[j])
            if i < len(filas) - 1:     # la tabla no se parte: cada fila se conserva con la siguiente
                t.cell(i, j).paragraphs[0].paragraph_format.keep_with_next = True
    return t


def emitir(doc, bl) -> None:
    for tipo, val in bl:
        if tipo == "h1":
            doc.add_paragraph(val, style="Heading 1").paragraph_format.keep_with_next = True
        elif tipo == "h2":
            doc.add_paragraph(val, style="Heading 2").paragraph_format.keep_with_next = True
        elif tipo == "p":
            p = doc.add_paragraph()
            _espaciado(p, 120, 120, jc=3)          # justificado
            _texto(p, val)
        elif tipo == "li":
            p = doc.add_paragraph()
            _espaciado(p, 60, 60, jc=3)
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Twips(360), Twips(-360)
            pf.tab_stops.add_tab_stop(Twips(360))
            p.add_run("•\t")
            _texto(p, val)
        elif tipo == "img":
            p = doc.add_paragraph()
            _espaciado(p, 120, 120, jc=1)          # centrado
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(RAIZ / val), width=Inches(ANCHO_FIG))
        elif tipo in ("fig", "tab"):
            p = doc.add_paragraph()
            _espaciado(p, 40, 40, jc=3 if tipo == "fig" else None)
            p.paragraph_format.keep_with_next = True
            _texto(p, val, tam=9, negrita=True)
        elif tipo == "fuente":
            p = doc.add_paragraph()
            _espaciado(p, 0, 160, jc=3)
            _texto(p, val, tam=8, cursiva=True, color=GRIS)
        elif tipo == "tabla":
            _tabla(doc, val)


# ---------------------------------------------------------------- documentos

def _sin_comentarios(doc) -> None:
    """Quita del documento solo las partes de comentarios (la plantilla no lleva el texto de JP)."""
    for rid, rel in list(doc.part.rels.items()):
        if any(k in rel.reltype for k in ("/comments", "commentsExtended", "commentsIds", "commentsExtensible", "/people")):
            del doc.part.rels[rid]


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def marcas(p: Path) -> dict:
    with zipfile.ZipFile(p) as z:
        doc = z.read("word/document.xml").decode("utf-8")
        com = z.read("word/comments.xml").decode("utf-8") if "word/comments.xml" in z.namelist() else ""
    return {"comentarios": len(re.findall(r"<w:comment ", com)),
            "referencias a comentarios": len(re.findall(r"<w:commentReference ", doc)),
            "inserciones": len(re.findall(r"<w:ins ", doc)),
            "eliminaciones": len(re.findall(r"<w:del ", doc)),
            "titulos 1": len(re.findall(r'w:pStyle w:val="Heading1"', doc)),
            "imagenes": len(re.findall(r"<pic:pic[ >]", doc)),
            "notas al pie": len(re.findall(r"<w:footnoteReference ", doc))}


def main() -> None:
    h0 = (RAIZ / "data/hash_original_antes.txt").read_text().split()[0]
    assert sha(ORIGINAL) == h0, "el original cambió antes de generar los .docx"
    bl = bloques(FUENTE_MD.read_text(encoding="utf-8"))

    # 1. la sección sola, sobre la plantilla de JP
    sola = Document(str(ORIGINAL))
    cuerpo = sola.element.body
    for hijo in list(cuerpo):
        if hijo.tag != qn("w:sectPr"):
            cuerpo.remove(hijo)
    _sin_comentarios(sola)
    for rid, rel in list(sola.part.rels.items()):          # imágenes de JP que ya no se usan
        if rel.reltype.endswith("/image"):
            del sola.part.rels[rid]
    NOTA_DOC[0] = _Notas(sola, limpiar=True)
    emitir(sola, bl)
    NOTA_DOC[0].guardar()
    sola.save(str(SALIDA_SOLA))

    # 2. la copia del documento de JP con la sección al final (opción b)
    copia = Document(str(ORIGINAL))
    copia.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    NOTA_DOC[0] = _Notas(copia, limpiar=False)
    emitir(copia, bl)
    NOTA_DOC[0].guardar()
    copia.save(str(SALIDA_COPIA))

    h1 = sha(ORIGINAL)
    (RAIZ / "data/hash_original_despues.txt").write_text(f"{h1}  {ORIGINAL.name}\n")
    antes, despues, sec = marcas(ORIGINAL), marcas(SALIDA_COPIA), marcas(SALIDA_SOLA)
    print("hash original antes  :", h0)
    print("hash original después:", h1, "(sin cambio)" if h1 == h0 else "(CAMBIÓ)")
    for k in antes:
        print(f"{k:28s} original {antes[k]:4d} · sección {sec[k]:4d} · copia {despues[k]:4d}")
    ok = h1 == h0 and all(despues[k] == antes[k] + sec[k] for k in antes) and sec["comentarios"] == 0
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
